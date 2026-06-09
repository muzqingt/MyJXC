from flask import render_template, redirect, url_for, flash, request, jsonify, current_app, abort
from flask_login import login_required, current_user
from app import db
from app.models import Product, Category, Log, PurchaseOrderItem, SalesOrderItem, StockInItem, StockOutItem, StockLog
from app.forms import ProductForm, CategoryForm, SearchForm
from datetime import datetime
from decimal import Decimal
import os
from werkzeug.utils import secure_filename
from flask import Blueprint
from sqlalchemy.exc import SQLAlchemyError, IntegrityError
from app.utils import to_decimal

def _save_product_image(file, product):
    if file.content_type not in ['image/jpeg', 'image/png', 'image/gif', 'image/jpg']:
        return '只能上传 JPG/PNG/GIF 格式图片'
    file.seek(0)
    header = file.read(8)
    file.seek(0)
    if not (header[:3] == b'\xff\xd8\xff' or header[:4] == b'\x89PNG' or header[:4] == b'GIF8'):
        return '只能上传 JPG/PNG/GIF 格式图片'
    file.seek(0, 2)
    size = file.tell()
    file.seek(0)
    if size > 2 * 1024 * 1024:
        return '图片大小不能超过 2MB'
    filename = secure_filename(file.filename)
    if filename:
        file_ext = os.path.splitext(filename)[1].lower()
        if file_ext in ['.jpg', '.jpeg', '.png', '.gif']:
            if product.image_path:
                old_filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], product.image_path)
                if os.path.exists(old_filepath):
                    os.remove(old_filepath)
            safe_code = secure_filename(product.code)
            if not safe_code:
                safe_code = f'product_{product.id}'
            new_filename = f"product_{safe_code}{file_ext}"
            filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], new_filename)
            file.save(filepath)
            product.image_path = new_filename
    return None


# 创建蓝图
bp = Blueprint('product', __name__, url_prefix='/product')

# 商品管理
@bp.route('/')
@bp.route('/products')
@login_required
def index():
    page = request.args.get('page', 1, type=int)
    per_page = 20
    
    search_form = SearchForm()
    keyword = request.args.get('keyword', '')
    category_id = request.args.get('category_id', type=int)
    
    query = Product.query
    
    if keyword:
        query = query.filter(
            Product.name.contains(keyword) | 
            Product.code.contains(keyword) |
            Product.specification.contains(keyword)
        )
    
    if category_id:
        query = query.filter_by(category_id=category_id)
    
    products = query.order_by(Product.updated_at.desc()).paginate(page=page, per_page=per_page, error_out=False)
    categories = Category.query.all()
    
    return render_template('product/index.html', 
                         title='商品管理',
                         products=products,
                         categories=categories,
                         search_form=search_form,
                         keyword=keyword,
                         category_id=category_id)

@bp.route('/products/new', methods=['GET', 'POST'])
@login_required
def new_product():
    form = ProductForm()
    form.category_id.choices = [(0, '无分类')] + [(c.id, c.name) for c in Category.query.all()]
    
    if form.validate_on_submit():
        # 检查编号是否已存在
        if Product.query.filter_by(code=form.code.data).first():
            flash('商品编号已存在，请使用其他编号！', 'danger')
            return render_template('product/edit.html', title='新建商品', form=form)
        
        product = Product(
            code=form.code.data,
            name=form.name.data,
            category_id=form.category_id.data if form.category_id.data != 0 else None,
            specification=form.specification.data,
            unit=form.unit.data,
            purchase_price=form.purchase_price.data or 0,
            sale_price=form.sale_price.data or 0,
            safety_stock=form.safety_stock.data or 0,
            description=form.description.data
        )
        
        if form.image.data:
            error = _save_product_image(form.image.data, product)
            if error:
                flash(error, 'danger')
                return redirect(url_for('product.new_product'))

        try:
            db.session.add(product)
            log = Log(
                user_id=current_user.id,
                action='添加商品',
                details=f'添加商品: {product.code} - {product.name}',
                ip_address=request.remote_addr
            )
            db.session.add(log)
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            if product.image_path:
                filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], product.image_path)
                if os.path.exists(filepath):
                    os.remove(filepath)
            flash('商品编号已存在，请使用其他编号！', 'danger')
            return redirect(url_for('product.new_product'))
        except SQLAlchemyError:
            db.session.rollback()
            if product.image_path:
                filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], product.image_path)
                if os.path.exists(filepath):
                    os.remove(filepath)
            flash('商品添加失败，请重试！', 'danger')
            return redirect(url_for('product.new_product'))

        flash('商品添加成功！', 'success')
        return redirect(url_for('product.index'))
    
    return render_template('product/edit.html', 
                         title='添加商品',
                         form=form,
                         action='new')

@bp.route('/products/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit_product(id):
    product = db.session.get(Product, id)
    if product is None:
        abort(404)
    form = ProductForm(obj=product)
    form.category_id.choices = [(0, '无分类')] + [(c.id, c.name) for c in Category.query.all()]
    
    if form.validate_on_submit():
        # 检查编号是否已存在（排除自身）
        existing = Product.query.filter_by(code=form.code.data).first()
        if existing and existing.id != product.id:
            flash('商品编号已存在，请使用其他编号！', 'danger')
            return render_template('product/edit.html', title='编辑商品', form=form, product=product)
        
        product.code = form.code.data
        product.name = form.name.data
        product.category_id = form.category_id.data if form.category_id.data != 0 else None
        product.specification = form.specification.data
        product.unit = form.unit.data
        product.purchase_price = form.purchase_price.data or 0
        product.sale_price = form.sale_price.data or 0
        product.safety_stock = form.safety_stock.data or 0
        product.description = form.description.data
        product.updated_at = datetime.now()
        
        if form.image.data:
            error = _save_product_image(form.image.data, product)
            if error:
                flash(error, 'danger')
                return redirect(url_for('product.edit_product', id=product.id))

        log = Log(
            user_id=current_user.id,
            action='修改商品',
            details=f'修改商品: {product.code} - {product.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            flash('商品修改失败，请重试！', 'danger')
            return redirect(url_for('product.edit_product', id=product.id))

        flash('商品修改成功！', 'success')
        return redirect(url_for('product.index'))
    
    return render_template('product/edit.html', 
                         title='编辑商品',
                         form=form,
                         product=product,
                         action='edit')

@bp.route('/products/<int:id>/delete', methods=['POST'])
@login_required
def delete_product(id):
    product = db.session.get(Product, id)
    if product is None:
        abort(404)

    # 检查是否有相关记录
    if (db.session.query(PurchaseOrderItem).filter_by(product_id=id).first() is not None
            or db.session.query(SalesOrderItem).filter_by(product_id=id).first() is not None
            or db.session.query(StockInItem).filter_by(product_id=id).first() is not None
            or db.session.query(StockOutItem).filter_by(product_id=id).first() is not None
            or db.session.query(StockLog).filter_by(product_id=id).first() is not None):
        flash('该商品已有库存或订单记录，无法删除！', 'danger')
        return redirect(url_for('product.index'))
    
    # 删除图片
    if product.image_path:
        filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], product.image_path)
        if os.path.exists(filepath):
            os.remove(filepath)
    
    # 记录日志
    log = Log(
        user_id=current_user.id,
        action='删除商品',
        details=f'删除商品: {product.code} - {product.name}',
        ip_address=request.remote_addr
    )
    db.session.add(log)

    try:
        db.session.delete(product)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash('删除商品失败，请重试！', 'danger')
        return redirect(url_for('product.index'))

    flash('商品删除成功！', 'success')
    return redirect(url_for('product.index'))

@bp.route('/products/<int:id>')
@login_required
def view_product(id):
    product = db.session.get(Product, id)
    if product is None:
        abort(404)
    return render_template('product/view.html',
                         title='商品详情',
                         product=product)

# 分类管理
@bp.route('/categories')
@login_required
def categories():
    categories = Category.query.all()
    return render_template('product/categories.html', 
                         title='商品分类',
                         categories=categories)

@bp.route('/categories/new', methods=['GET', 'POST'])
@login_required
def new_category():
    form = CategoryForm()
    form.parent_id.choices = [(0, '无父级')] + [(c.id, c.name) for c in Category.query.all()]
    
    if form.validate_on_submit():
        category = Category(
            name=form.name.data,
            parent_id=form.parent_id.data if form.parent_id.data != 0 else None,
            description=form.description.data
        )
        
        try:
            db.session.add(category)
            db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            flash('分类添加失败，请重试！', 'danger')
            return redirect(url_for('product.new_category'))

        flash('分类添加成功！', 'success')
        return redirect(url_for('product.categories'))
    
    return render_template('product/category_edit.html', 
                         title='添加分类',
                         form=form,
                         action='new')

@bp.route('/categories/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit_category(id):
    category = db.session.get(Category, id)
    if category is None:
        abort(404)
    form = CategoryForm(obj=category)
    form.parent_id.choices = [(0, '无父级')] + [(c.id, c.name) for c in Category.query.filter(Category.id != id).all()]
    
    if form.validate_on_submit():
        category.name = form.name.data
        category.parent_id = form.parent_id.data if form.parent_id.data != 0 else None
        category.description = form.description.data

        try:
            db.session.commit()
        except SQLAlchemyError:
            db.session.rollback()
            flash('分类修改失败，请重试！', 'danger')
            return redirect(url_for('product.edit_category', id=category.id))

        flash('分类修改成功！', 'success')
        return redirect(url_for('product.categories'))
    
    return render_template('product/category_edit.html', 
                         title='编辑分类',
                         form=form,
                         category=category,
                         action='edit')

@bp.route('/categories/<int:id>/delete', methods=['POST'])
@login_required
def delete_category(id):
    category = db.session.get(Category, id)
    if category is None:
        abort(404)

    # 检查是否有子分类
    if db.session.query(Category).filter_by(parent_id=id).first() is not None:
        flash('该分类下有子分类，无法删除！', 'danger')
        return redirect(url_for('product.categories'))

    # 检查是否有商品
    if db.session.query(Product).filter_by(category_id=id).first() is not None:
        flash('该分类下有商品，无法删除！', 'danger')
        return redirect(url_for('product.categories'))
    
    try:
        db.session.delete(category)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash('分类删除失败，请重试！', 'danger')
        return redirect(url_for('product.categories'))

    flash('分类删除成功！', 'success')
    return redirect(url_for('product.categories'))

# API接口
@bp.route('/api/products')
@login_required
def api_products():
    limit = request.args.get('limit', 100, type=int)
    limit = min(max(limit, 1), 1000)
    products = Product.query.limit(limit).all()
    result = []
    for product in products:
        result.append({
            'id': product.id,
            'code': product.code,
            'name': product.name,
            'specification': product.specification,
            'unit': product.unit,
            'purchase_price': str(product.purchase_price),
            'sale_price': str(product.sale_price),
            'stock_quantity': str(product.stock_quantity)
        })
    return jsonify(result)

@bp.route('/api/product/price', methods=['POST'])
@login_required
def api_update_price():
    """更新商品价格"""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'message': '无效的JSON数据'})

    product_id = data.get('product_id')
    price_type = data.get('price_type')  # 'sale_price' or 'purchase_price'
    price = data.get('price')

    if not product_id or not price_type or price is None:
        return jsonify({'success': False, 'message': '参数不完整'})
    
    try:
        product_id = int(product_id)
    except (ValueError, TypeError):
        return jsonify({'success': False, 'message': '商品ID无效'})
    
    product = db.session.get(Product, product_id)
    if not product:
        return jsonify({'success': False, 'message': '商品不存在'})
    
    try:
        price_value = to_decimal(price)
        if price_value < 0:
            return jsonify({'success': False, 'message': '价格不能为负数'})
        if price_type == 'sale_price':
            product.sale_price = price_value
        elif price_type == 'purchase_price':
            product.purchase_price = price_value
        else:
            return jsonify({'success': False, 'message': '价格类型无效'})
        
        db.session.commit()
        return jsonify({'success': True, 'message': '价格已更新'})
    except SQLAlchemyError as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': '数据库错误，请重试'})

@bp.route('/api/categories')
@login_required
def api_categories():
    categories = Category.query.all()
    result = []
    for category in categories:
        result.append({
            'id': category.id,
            'name': category.name,
            'parent_id': category.parent_id
        })
    return jsonify(result)


@bp.route('/import')
@login_required
def import_page():
    """商品导入页面"""
    return render_template('product/import.html', title='商品导入')


@bp.route('/import/template')
@login_required
def import_template():
    """下载商品导入模板"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from flask import send_file
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = '商品导入模板'

    # 表头
    headers = ['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    # 示例数据
    ws.append(['SKU001', '示例商品', '默认分类', '个', 10.00, 20.00, 10])

    # 调整列宽
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='商品导入模板.xlsx'
    )


@bp.route('/import', methods=['POST'])
@login_required
def import_products():
    """处理商品导入"""
    from openpyxl import load_workbook
    import io

    if 'file' not in request.files:
        flash('请选择文件！', 'danger')
        return redirect(url_for('product.import_page'))

    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        flash('请上传 .xlsx 格式文件！', 'danger')
        return redirect(url_for('product.import_page'))

    try:
        wb = load_workbook(io.BytesIO(file.read()))
        ws = wb.active
    except Exception as e:
        flash('文件格式错误！', 'danger')
        return redirect(url_for('product.import_page'))

    # 获取分类映射
    categories = {c.name: c.id for c in Category.query.all()}

    success_count = 0
    fail_count = 0
    skip_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue

        code = str(row[0]).strip()
        name = str(row[1]).strip() if row[1] else ''
        category_name = str(row[2]).strip() if row[2] else ''
        unit = str(row[3]).strip() if row[3] else '个'
        purchase_price = row[4] if row[4] else 0
        sale_price = row[5] if row[5] else 0
        safety_stock = row[6] if row[6] else 0

        # 校验
        if not code:
            errors.append(f'第 {row_idx} 行: 编码不能为空')
            fail_count += 1
            continue

        if not name:
            errors.append(f'第 {row_idx} 行: 名称不能为空')
            fail_count += 1
            continue

        # 检查编码唯一性
        if Product.query.filter_by(code=code).first():
            errors.append(f'第 {row_idx} 行: 编码 {code} 已存在')
            skip_count += 1
            continue

        # 获取分类 ID
        category_id = categories.get(category_name)
        if category_name and not category_id:
            errors.append(f'第 {row_idx} 行: 分类 {category_name} 不存在')
            fail_count += 1
            continue

        # 创建商品
        try:
            product = Product(
                code=code,
                name=name,
                category_id=category_id,
                unit=unit,
                purchase_price=Decimal(str(purchase_price)),
                sale_price=Decimal(str(sale_price)),
                safety_stock=Decimal(str(safety_stock)),
            )
            db.session.add(product)
            success_count += 1
        except Exception:
            errors.append(f'第 {row_idx} 行: 创建失败，数据格式有误')
            fail_count += 1

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash('导入失败，请重试！', 'danger')
        return redirect(url_for('product.import_page'))

    return render_template('product/import_result.html',
                         title='导入结果',
                         success_count=success_count,
                         fail_count=fail_count,
                         skip_count=skip_count,
                         errors=errors)