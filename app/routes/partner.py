from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import login_required, current_user
from app import db
from app.models import Supplier, Customer, Warehouse, PurchaseOrder, Payment, SalesOrder, Receipt, StockIn, StockOut, StockLog, Log
from app.forms import SupplierForm, CustomerForm, WarehouseForm
from flask import Blueprint
from sqlalchemy.exc import SQLAlchemyError, IntegrityError

# 创建蓝图
bp = Blueprint('partner', __name__, url_prefix='/partner')

# 供应商管理
@bp.route('/suppliers')
@login_required
def suppliers():
    suppliers = Supplier.query.order_by(Supplier.created_at.desc()).limit(500).all()
    return render_template('product/suppliers.html',
                         title='供应商管理',
                         suppliers=suppliers)

@bp.route('/suppliers/new', methods=['GET', 'POST'])
@login_required
def add_supplier():
    form = SupplierForm()

    if form.validate_on_submit():
        # 检查编号是否已存在
        if Supplier.query.filter_by(code=form.code.data).first():
            flash('供应商编号已存在，请使用其他编号！', 'danger')
            return render_template('product/supplier_edit.html', title='添加供应商', form=form, action='new')

        supplier = Supplier(
            code=form.code.data,
            name=form.name.data,
            contact_person=form.contact_person.data,
            phone=form.phone.data,
            address=form.address.data,
            email=form.email.data
        )

        db.session.add(supplier)
        log = Log(
            user_id=current_user.id,
            action='添加供应商',
            details=f'添加供应商: {supplier.code} - {supplier.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/supplier_edit.html', title='添加供应商', form=form, action='new')
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/supplier_edit.html', title='添加供应商', form=form, action='new')

        flash('供应商添加成功！', 'success')
        return redirect(url_for('partner.suppliers'))

    return render_template('product/supplier_edit.html',
                         title='添加供应商',
                         form=form,
                         action='new')

@bp.route('/suppliers/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit_supplier(id):
    supplier = db.session.get(Supplier, id)
    if supplier is None:
        abort(404)
    form = SupplierForm(obj=supplier)

    if form.validate_on_submit():
        # 检查编号是否已存在（排除自身）
        existing = Supplier.query.filter_by(code=form.code.data).first()
        if existing and existing.id != supplier.id:
            flash('供应商编号已存在，请使用其他编号！', 'danger')
            return render_template('product/supplier_edit.html', title='编辑供应商', form=form, supplier=supplier)

        supplier.code = form.code.data
        supplier.name = form.name.data
        supplier.contact_person = form.contact_person.data
        supplier.phone = form.phone.data
        supplier.address = form.address.data
        supplier.email = form.email.data

        log = Log(
            user_id=current_user.id,
            action='修改供应商',
            details=f'修改供应商: {supplier.code} - {supplier.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/supplier_edit.html', title='编辑供应商', form=form, supplier=supplier)
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/supplier_edit.html', title='编辑供应商', form=form, supplier=supplier)

        flash('供应商修改成功！', 'success')
        return redirect(url_for('partner.suppliers'))

    return render_template('product/supplier_edit.html',
                         title='编辑供应商',
                         form=form,
                         supplier=supplier,
                         action='edit')

@bp.route('/suppliers/<int:id>/delete', methods=['POST'])
@login_required
def delete_supplier(id):
    supplier = db.session.get(Supplier, id)
    if supplier is None:
        abort(404)

    has_orders = db.session.query(PurchaseOrder).filter_by(supplier_id=supplier.id).first() is not None
    has_payments = db.session.query(Payment).filter_by(supplier_id=supplier.id).first() is not None
    has_stock_ins = db.session.query(StockIn).join(PurchaseOrder, StockIn.purchase_order_id == PurchaseOrder.id).filter(PurchaseOrder.supplier_id == supplier.id).first() is not None
    has_stock_logs = db.session.query(StockLog).filter(StockLog.reference_type == 'purchase_order', StockLog.reference_id.in_(db.session.query(PurchaseOrder.id).filter_by(supplier_id=supplier.id))).first() is not None
    if has_orders or has_payments or has_stock_ins or has_stock_logs:
        flash('该供应商已有采购或付款记录，无法删除！', 'danger')
        return redirect(url_for('partner.suppliers'))

    log = Log(
        user_id=current_user.id,
        action='删除供应商',
        details=f'删除供应商: {supplier.code} - {supplier.name}',
        ip_address=request.remote_addr
    )
    db.session.add(log)
    try:
        db.session.delete(supplier)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash('数据库错误，删除失败！', 'danger')
        return redirect(url_for('partner.suppliers'))

    flash('供应商删除成功！', 'success')
    return redirect(url_for('partner.suppliers'))

# 客户管理
@bp.route('/customers')
@login_required
def customers():
    customers = Customer.query.order_by(Customer.created_at.desc()).limit(500).all()
    return render_template('product/customers.html',
                         title='客户管理',
                         customers=customers)

@bp.route('/customers/new', methods=['GET', 'POST'])
@login_required
def add_customer():
    form = CustomerForm()

    if form.validate_on_submit():
        # 检查编号是否已存在
        if Customer.query.filter_by(code=form.code.data).first():
            flash('客户编号已存在，请使用其他编号！', 'danger')
            return render_template('product/customer_edit.html', title='添加客户', form=form, action='new')

        customer = Customer(
            code=form.code.data,
            name=form.name.data,
            contact_person=form.contact_person.data,
            phone=form.phone.data,
            address=form.address.data,
            email=form.email.data
        )

        db.session.add(customer)
        log = Log(
            user_id=current_user.id,
            action='添加客户',
            details=f'添加客户: {customer.code} - {customer.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/customer_edit.html', title='添加客户', form=form, action='new')
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/customer_edit.html', title='添加客户', form=form, action='new')

        flash('客户添加成功！', 'success')
        return redirect(url_for('partner.customers'))

    return render_template('product/customer_edit.html',
                         title='添加客户',
                         form=form,
                         action='new')

@bp.route('/customers/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit_customer(id):
    customer = db.session.get(Customer, id)
    if customer is None:
        abort(404)
    form = CustomerForm(obj=customer)

    if form.validate_on_submit():
        # 检查编号是否已存在（排除自身）
        existing = Customer.query.filter_by(code=form.code.data).first()
        if existing and existing.id != customer.id:
            flash('客户编号已存在，请使用其他编号！', 'danger')
            return render_template('product/customer_edit.html', title='编辑客户', form=form, customer=customer)

        customer.code = form.code.data
        customer.name = form.name.data
        customer.contact_person = form.contact_person.data
        customer.phone = form.phone.data
        customer.address = form.address.data
        customer.email = form.email.data

        log = Log(
            user_id=current_user.id,
            action='修改客户',
            details=f'修改客户: {customer.code} - {customer.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/customer_edit.html', title='编辑客户', form=form, customer=customer)
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/customer_edit.html', title='编辑客户', form=form, customer=customer)

        flash('客户修改成功！', 'success')
        return redirect(url_for('partner.customers'))

    return render_template('product/customer_edit.html',
                         title='编辑客户',
                         form=form,
                         customer=customer,
                         action='edit')

@bp.route('/customers/<int:id>/delete', methods=['POST'])
@login_required
def delete_customer(id):
    customer = db.session.get(Customer, id)
    if customer is None:
        abort(404)

    # 检查是否有销售记录或收款记录
    has_sales = db.session.query(SalesOrder).filter_by(customer_id=customer.id).first() is not None
    has_receipts = db.session.query(Receipt).filter_by(customer_id=customer.id).first() is not None
    if has_sales or has_receipts:
        flash('该客户已有销售或收款记录，无法删除！', 'danger')
        return redirect(url_for('partner.customers'))

    log = Log(
        user_id=current_user.id,
        action='删除客户',
        details=f'删除客户: {customer.code} - {customer.name}',
        ip_address=request.remote_addr
    )
    db.session.add(log)
    try:
        db.session.delete(customer)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash('数据库错误，删除失败！', 'danger')
        return redirect(url_for('partner.customers'))

    flash('客户删除成功！', 'success')
    return redirect(url_for('partner.customers'))

# 仓库管理
@bp.route('/warehouses')
@login_required
def warehouses():
    warehouses = Warehouse.query.order_by(Warehouse.created_at.desc()).limit(500).all()
    return render_template('product/warehouses.html',
                         title='仓库管理',
                         warehouses=warehouses)

@bp.route('/warehouses/new', methods=['GET', 'POST'])
@login_required
def add_warehouse():
    form = WarehouseForm()

    if form.validate_on_submit():
        # 检查编号是否已存在
        if Warehouse.query.filter_by(code=form.code.data).first():
            flash('仓库编号已存在，请使用其他编号！', 'danger')
            return render_template('product/warehouse_edit.html', title='添加仓库', form=form, action='new')

        warehouse = Warehouse(
            code=form.code.data,
            name=form.name.data,
            address=form.address.data,
            manager=form.manager.data,
            phone=form.phone.data
        )

        db.session.add(warehouse)
        log = Log(
            user_id=current_user.id,
            action='添加仓库',
            details=f'添加仓库: {warehouse.code} - {warehouse.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/warehouse_edit.html', title='添加仓库', form=form, action='new')
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/warehouse_edit.html', title='添加仓库', form=form, action='new')

        flash('仓库添加成功！', 'success')
        return redirect(url_for('partner.warehouses'))

    return render_template('product/warehouse_edit.html',
                         title='添加仓库',
                         form=form,
                         action='new')

@bp.route('/warehouses/<int:id>/edit', methods=['GET', 'POST'])
@login_required
def edit_warehouse(id):
    warehouse = db.session.get(Warehouse, id)
    if warehouse is None:
        abort(404)
    form = WarehouseForm(obj=warehouse)

    if form.validate_on_submit():
        # 检查编号是否已存在（排除自身）
        existing = Warehouse.query.filter_by(code=form.code.data).first()
        if existing and existing.id != warehouse.id:
            flash('仓库编号已存在，请使用其他编号！', 'danger')
            return render_template('product/warehouse_edit.html', title='编辑仓库', form=form, warehouse=warehouse)

        warehouse.code = form.code.data
        warehouse.name = form.name.data
        warehouse.address = form.address.data
        warehouse.manager = form.manager.data
        warehouse.phone = form.phone.data

        log = Log(
            user_id=current_user.id,
            action='修改仓库',
            details=f'修改仓库: {warehouse.code} - {warehouse.name}',
            ip_address=request.remote_addr
        )
        db.session.add(log)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            flash('编码已存在，请使用其他编码！', 'danger')
            return render_template('product/warehouse_edit.html', title='编辑仓库', form=form, warehouse=warehouse)
        except SQLAlchemyError:
            db.session.rollback()
            flash('数据库错误，操作失败！', 'danger')
            return render_template('product/warehouse_edit.html', title='编辑仓库', form=form, warehouse=warehouse)

        flash('仓库修改成功！', 'success')
        return redirect(url_for('partner.warehouses'))

    return render_template('product/warehouse_edit.html',
                         title='编辑仓库',
                         form=form,
                         warehouse=warehouse,
                         action='edit')

@bp.route('/warehouses/<int:id>/delete', methods=['POST'])
@login_required
def delete_warehouse(id):
    warehouse = db.session.get(Warehouse, id)
    if warehouse is None:
        abort(404)

    # 检查是否有库存记录
    has_stock_ins = db.session.query(StockIn).filter_by(warehouse_id=warehouse.id).first() is not None
    has_stock_outs = db.session.query(StockOut).filter_by(warehouse_id=warehouse.id).first() is not None
    has_stock_logs = db.session.query(StockLog).filter_by(warehouse_id=warehouse.id).first() is not None
    if has_stock_ins or has_stock_outs or has_stock_logs:
        flash('该仓库已有库存记录，无法删除！', 'danger')
        return redirect(url_for('partner.warehouses'))

    log = Log(
        user_id=current_user.id,
        action='删除仓库',
        details=f'删除仓库: {warehouse.code} - {warehouse.name}',
        ip_address=request.remote_addr
    )
    db.session.add(log)
    try:
        db.session.delete(warehouse)
        db.session.commit()
    except SQLAlchemyError:
        db.session.rollback()
        flash('数据库错误，删除失败！', 'danger')
        return redirect(url_for('partner.warehouses'))

    flash('仓库删除成功！', 'success')
    return redirect(url_for('partner.warehouses'))


@bp.route('/suppliers/import')
@login_required
def supplier_import_page():
    """供应商导入页面"""
    return render_template('partner/supplier_import.html', title='供应商导入')


@bp.route('/suppliers/import/template')
@login_required
def supplier_import_template():
    """下载供应商导入模板"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from flask import send_file
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = '供应商导入模板'

    headers = ['供应商编码', '供应商名称', '联系人', '电话', '地址', '邮箱']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    ws.append(['SUP001', '示例供应商', '张三', '13800000000', '示例地址', 'example@test.com'])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='供应商导入模板.xlsx'
    )


@bp.route('/suppliers/import', methods=['POST'])
@login_required
def import_suppliers():
    """处理供应商导入"""
    from openpyxl import load_workbook
    import io
    from decimal import Decimal

    if 'file' not in request.files:
        flash('请选择文件！', 'danger')
        return redirect(url_for('partner.supplier_import_page'))

    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        flash('请上传 .xlsx 格式文件！', 'danger')
        return redirect(url_for('partner.supplier_import_page'))

    try:
        wb = load_workbook(io.BytesIO(file.read()))
        ws = wb.active
    except Exception:
        flash('文件格式错误！', 'danger')
        return redirect(url_for('partner.supplier_import_page'))

    success_count = 0
    fail_count = 0
    skip_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue

        code = str(row[0]).strip()
        name = str(row[1]).strip() if row[1] else ''
        contact = str(row[2]).strip() if row[2] else ''
        phone = str(row[3]).strip() if row[3] else ''
        address = str(row[4]).strip() if row[4] else ''
        email = str(row[5]).strip() if row[5] else ''

        if not code:
            errors.append(f'第 {row_idx} 行: 编码不能为空')
            fail_count += 1
            continue

        if not name:
            errors.append(f'第 {row_idx} 行: 名称不能为空')
            fail_count += 1
            continue

        if Supplier.query.filter_by(code=code).first():
            errors.append(f'第 {row_idx} 行: 编码 {code} 已存在')
            skip_count += 1
            continue

        try:
            supplier = Supplier(
                code=code,
                name=name,
                contact_person=contact,
                phone=phone,
                address=address,
                email=email,
            )
            db.session.add(supplier)
            success_count += 1
        except Exception as e:
            errors.append(f'第 {row_idx} 行: 创建失败 - {str(e)}')
            fail_count += 1

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash('导入失败，请重试！', 'danger')
        return redirect(url_for('partner.supplier_import_page'))

    return render_template('partner/import_result.html',
                         title='导入结果',
                         success_count=success_count,
                         fail_count=fail_count,
                         skip_count=skip_count,
                         errors=errors,
                         back_url=url_for('partner.supplier_import_page'),
                         list_url=url_for('partner.suppliers'))


@bp.route('/customers/import')
@login_required
def customer_import_page():
    """客户导入页面"""
    return render_template('partner/customer_import.html', title='客户导入')


@bp.route('/customers/import/template')
@login_required
def customer_import_template():
    """下载客户导入模板"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from flask import send_file
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = '客户导入模板'

    headers = ['客户编码', '客户名称', '联系人', '电话', '地址', '邮箱']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    ws.append(['CUS001', '示例客户', '李四', '13900000000', '示例地址', 'example@test.com'])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='客户导入模板.xlsx'
    )


@bp.route('/customers/import', methods=['POST'])
@login_required
def import_customers():
    """处理客户导入"""
    from openpyxl import load_workbook
    import io
    from decimal import Decimal

    if 'file' not in request.files:
        flash('请选择文件！', 'danger')
        return redirect(url_for('partner.customer_import_page'))

    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        flash('请上传 .xlsx 格式文件！', 'danger')
        return redirect(url_for('partner.customer_import_page'))

    try:
        wb = load_workbook(io.BytesIO(file.read()))
        ws = wb.active
    except Exception:
        flash('文件格式错误！', 'danger')
        return redirect(url_for('partner.customer_import_page'))

    success_count = 0
    fail_count = 0
    skip_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue

        code = str(row[0]).strip()
        name = str(row[1]).strip() if row[1] else ''
        contact = str(row[2]).strip() if row[2] else ''
        phone = str(row[3]).strip() if row[3] else ''
        address = str(row[4]).strip() if row[4] else ''
        email = str(row[5]).strip() if row[5] else ''

        if not code:
            errors.append(f'第 {row_idx} 行: 编码不能为空')
            fail_count += 1
            continue

        if not name:
            errors.append(f'第 {row_idx} 行: 名称不能为空')
            fail_count += 1
            continue

        if Customer.query.filter_by(code=code).first():
            errors.append(f'第 {row_idx} 行: 编码 {code} 已存在')
            skip_count += 1
            continue

        try:
            customer = Customer(
                code=code,
                name=name,
                contact_person=contact,
                phone=phone,
                address=address,
                email=email,
            )
            db.session.add(customer)
            success_count += 1
        except Exception as e:
            errors.append(f'第 {row_idx} 行: 创建失败 - {str(e)}')
            fail_count += 1

    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash('导入失败，请重试！', 'danger')
        return redirect(url_for('partner.customer_import_page'))

    return render_template('partner/import_result.html',
                         title='导入结果',
                         success_count=success_count,
                         fail_count=fail_count,
                         skip_count=skip_count,
                         errors=errors,
                         back_url=url_for('partner.customer_import_page'),
                         list_url=url_for('partner.customers'))
