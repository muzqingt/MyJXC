"""
完整单元测试 — 覆盖所有缺失的代码路径
"""
import pytest
from decimal import Decimal
from datetime import datetime, date
from app import db
from app.models import (
    User, Product, Category, Supplier, Customer, Warehouse,
    PurchaseOrder, PurchaseOrderItem, SalesOrder, SalesOrderItem,
    StockIn, StockInItem, StockOut, StockOutItem,
    PurchaseReturn, PurchaseReturnItem, SalesReturn, SalesReturnItem,
    StockLog, SystemSetting, Receipt, Payment, Expense, Log
)
from app.utils import (
    to_decimal, add_balance, sub_balance, generate_order_number,
    generate_order_number_safe, build_products_data, build_partners_data,
    flash_success, flash_error, flash_warning, safe_commit,
    localize_dt, format_local_dt, get_redirect_tab,
    apply_excel_header_style, EXCEL_HEADER_FONT, EXCEL_HEADER_FILL,
    EXCEL_THIN_BORDER, EXCEL_HEADER_ALIGNMENT
)
from app.utils_order import create_stock_log, match_reference_id_filter, rebuild_items_on_error
from app.constants import OrderStatus, ChangeType, PaymentMethod, UserRole


# ==================== to_decimal 测试 ====================

def test_to_decimal_none():
    """None 值返回默认值"""
    assert to_decimal(None) == Decimal('0')
    assert to_decimal(None, Decimal('10')) == Decimal('10')


def test_to_decimal_decimal():
    """Decimal 值直接返回"""
    assert to_decimal(Decimal('123.45')) == Decimal('123.45')


def test_to_decimal_int():
    """整数转换"""
    assert to_decimal(100) == Decimal('100')


def test_to_decimal_float():
    """浮点数转换"""
    assert to_decimal(3.14) == Decimal('3.14')


def test_to_decimal_string():
    """字符串转换"""
    assert to_decimal('99.99') == Decimal('99.99')


def test_to_decimal_invalid_string():
    """无效字符串返回默认值"""
    assert to_decimal('abc') == Decimal('0')
    assert to_decimal('abc', Decimal('5')) == Decimal('5')


def test_to_decimal_other_type():
    """其他类型转换"""
    # 布尔值通过 int 转换
    assert to_decimal(int(True)) == Decimal('1')
    assert to_decimal(int(False)) == Decimal('0')


# ==================== add_balance / sub_balance 测试 ====================

def test_add_balance_customer(app, db_session):
    """增加客户余额"""
    customer = Customer(code='BAL_CUS', name='余额测试客户', contact_person='测试', phone='13800000000')
    db.session.add(customer)
    db.session.commit()

    add_balance(customer, 'receivable_balance', 100)
    assert customer.receivable_balance == Decimal('100')


def test_sub_balance_customer(app, db_session):
    """减少客户余额"""
    customer = Customer(code='BAL_CUS2', name='余额测试客户2', contact_person='测试', phone='13800000000')
    db.session.add(customer)
    db.session.commit()

    add_balance(customer, 'receivable_balance', 200)
    sub_balance(customer, 'receivable_balance', 50)
    assert customer.receivable_balance == Decimal('150')


def test_add_balance_supplier(app, db_session):
    """增加供应商余额"""
    supplier = Supplier(code='BAL_SUP', name='余额测试供应商', contact_person='测试', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()

    add_balance(supplier, 'payable_balance', 100)
    assert supplier.payable_balance == Decimal('100')


def test_sub_balance_supplier(app, db_session):
    """减少供应商余额"""
    supplier = Supplier(code='BAL_SUP2', name='余额测试供应商2', contact_person='测试', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()

    add_balance(supplier, 'payable_balance', 200)
    sub_balance(supplier, 'payable_balance', 50)
    assert supplier.payable_balance == Decimal('150')


# ==================== generate_order_number 测试 ====================

def test_generate_order_number_first(app, db_session):
    """生成第一个订单号"""
    prefix = 'PO'
    order_number = generate_order_number(prefix, PurchaseOrder)
    today = datetime.now().strftime('%Y%m%d')
    assert order_number == f'{prefix}{today}0001'


def test_generate_order_number_increment(app, db_session):
    """订单号自增"""
    supplier = Supplier(code='ORD_SUP', name='订单测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='ORD_WH', name='订单测试仓库')
    db.session.add_all([supplier, warehouse])
    db.session.commit()

    # 创建第一个订单
    order1 = PurchaseOrder(
        order_number=generate_order_number('PO', PurchaseOrder),
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order1)
    db.session.commit()

    # 生成第二个订单号
    order_number2 = generate_order_number('PO', PurchaseOrder)
    today = datetime.now().strftime('%Y%m%d')
    assert order_number2 == f'PO{today}0002'


def test_generate_order_number_different_prefix(app, db_session):
    """不同前缀的订单号"""
    so_number = generate_order_number('SO', SalesOrder)
    si_number = generate_order_number('SI', StockIn)
    today = datetime.now().strftime('%Y%m%d')
    assert so_number == f'SO{today}0001'
    assert si_number == f'SI{today}0001'


def test_generate_order_number_safe(app, db_session):
    """安全生成订单号"""
    order_number = generate_order_number_safe('PO', PurchaseOrder)
    assert order_number is not None
    assert 'PO' in order_number


# ==================== build_products_data 测试 ====================

def test_build_products_data_basic(app, db_session):
    """构建商品数据 - 基本"""
    product = Product(code='BPD_PROD', name='测试商品', unit='个',
                     purchase_price=50, sale_price=100, stock_quantity=10)
    db.session.add(product)
    db.session.commit()

    result = build_products_data([product])
    assert len(result) == 1
    assert result[0]['id'] == product.id
    assert result[0]['code'] == 'BPD_PROD'
    assert result[0]['name'] == '测试商品'


def test_build_products_data_exclude_prices(app, db_session):
    """构建商品数据 - 排除价格"""
    product = Product(code='BPD_PROD2', name='测试商品2', unit='个',
                     purchase_price=50, sale_price=100)
    db.session.add(product)
    db.session.commit()

    result = build_products_data([product], include_purchase_price=False, include_sale_price=False)
    assert 'purchase_price' not in result[0]
    assert 'sale_price' not in result[0]


def test_build_partners_data(app, db_session):
    """构建合作伙伴数据"""
    supplier = Supplier(code='BPD_SUP', name='测试供应商', contact_person='测试', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()

    result = build_partners_data([supplier])
    assert len(result) == 1
    assert result[0]['id'] == supplier.id
    assert result[0]['code'] == 'BPD_SUP'


# ==================== flash 函数测试 ====================

def test_flash_success_no_punctuation(app):
    """成功消息 - 无标点"""
    with app.test_request_context():
        result = flash_success('操作成功')
        assert result is None


def test_flash_success_with_punctuation(app):
    """成功消息 - 有标点"""
    with app.test_request_context():
        result = flash_success('操作成功！')
        assert result is None


def test_flash_error_no_punctuation(app):
    """错误消息 - 无标点"""
    with app.test_request_context():
        result = flash_error('操作失败')
        assert result is None


def test_flash_warning_no_punctuation(app):
    """警告消息 - 无标点"""
    with app.test_request_context():
        result = flash_warning('警告信息')
        assert result is None


# ==================== safe_commit 测试 ====================

def test_safe_commit_success(app, db_session):
    """安全提交 - 成功"""
    product = Product(code='SC_PROD', name='安全提交测试', unit='个')
    db.session.add(product)
    success, error = safe_commit()
    assert success is True
    assert error is None


def test_generate_order_number_no_filter_field(app, db_session):
    """生成订单号 - 无过滤字段"""
    # 测试没有 order_number/receipt_number 等字段的模型
    # 这会触发 last_order = None 分支
    order_number = generate_order_number('TEST', Log)
    today = datetime.now().strftime('%Y%m%d')
    assert order_number == f'TEST{today}0001'


def test_generate_order_number_with_existing(app, db_session):
    """生成订单号 - 已有订单"""
    supplier = Supplier(code='GON_SUP', name='订单号测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='GON_WH', name='订单号测试仓库')
    db.session.add_all([supplier, warehouse])
    db.session.commit()

    # 创建已有订单
    order = PurchaseOrder(
        order_number='PO202401010005',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    # 生成新订单号
    order_number = generate_order_number('PO', PurchaseOrder)
    today = datetime.now().strftime('%Y%m%d')
    # 应该生成今天的订单号，不是2024年的
    assert today in order_number


def test_generate_order_number_safe_no_filter(app, db_session):
    """安全生成订单号 - 无过滤字段"""
    order_number = generate_order_number_safe('TEST', Log)
    assert order_number is not None
    assert 'TEST' in order_number


def test_update_stock_and_log_in(app, db_session):
    """更新库存并记录日志 - 入库"""
    product = Product(code='USL_PROD', name='库存更新测试', unit='个', stock_quantity=Decimal('0'))
    warehouse = Warehouse(code='USL_WH', name='库存更新测试仓库')
    db.session.add_all([product, warehouse])
    db.session.commit()

    from app.utils import update_stock_and_log
    before, after = update_stock_and_log(
        product=product,
        warehouse_id=warehouse.id,
        quantity=Decimal('10'),
        change_type='in',
        reference_id=1,
        reference_type='stock_in',
        notes='测试入库',
        user_id=1
    )

    assert before == Decimal('0')
    assert after == Decimal('10')
    assert product.stock_quantity == Decimal('10')


def test_update_stock_and_log_out(app, db_session):
    """更新库存并记录日志 - 出库"""
    product = Product(code='USL_PROD2', name='库存更新测试2', unit='个', stock_quantity=Decimal('20'))
    warehouse = Warehouse(code='USL_WH2', name='库存更新测试仓库2')
    db.session.add_all([product, warehouse])
    db.session.commit()

    from app.utils import update_stock_and_log
    before, after = update_stock_and_log(
        product=product,
        warehouse_id=warehouse.id,
        quantity=Decimal('5'),
        change_type='out',
        reference_id=1,
        reference_type='stock_out',
        notes='测试出库',
        user_id=1
    )

    assert before == Decimal('20')
    assert after == Decimal('15')
    assert product.stock_quantity == Decimal('15')


def test_update_stock_and_log_invalid_type(app, db_session):
    """更新库存并记录日志 - 无效类型"""
    product = Product(code='USL_PROD3', name='库存更新测试3', unit='个', stock_quantity=Decimal('10'))
    warehouse = Warehouse(code='USL_WH3', name='库存更新测试仓库3')
    db.session.add_all([product, warehouse])
    db.session.commit()

    from app.utils import update_stock_and_log
    with pytest.raises(ValueError):
        update_stock_and_log(
            product=product,
            warehouse_id=warehouse.id,
            quantity=Decimal('5'),
            change_type='invalid_type',
            reference_id=1,
            reference_type='test',
            notes='测试',
            user_id=1
        )


def test_update_stock_and_log_negative_quantity(app, db_session):
    """更新库存并记录日志 - 负数数量"""
    product = Product(code='USL_PROD4', name='库存更新测试4', unit='个', stock_quantity=Decimal('10'))
    warehouse = Warehouse(code='USL_WH4', name='库存更新测试仓库4')
    db.session.add_all([product, warehouse])
    db.session.commit()

    from app.utils import update_stock_and_log
    with pytest.raises(ValueError):
        update_stock_and_log(
            product=product,
            warehouse_id=warehouse.id,
            quantity=Decimal('-5'),
            change_type='in',
            reference_id=1,
            reference_type='stock_in',
            notes='测试',
            user_id=1
        )


def test_apply_excel_header_style(app):
    """应用 Excel 表头样式"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active

    # 添加表头
    ws.cell(row=1, column=1, value='列1')
    ws.cell(row=1, column=2, value='列2')
    ws.cell(row=1, column=3, value='列3')

    from app.utils import apply_excel_header_style
    apply_excel_header_style(ws, 1, 3)

    # 验证样式已应用
    cell = ws.cell(row=1, column=1)
    assert cell.font.bold is True


def test_flash_success_with_endpoint(app):
    """成功消息 - 带重定向"""
    with app.test_request_context('/test'):
        result = flash_success('操作成功', 'main.index')
        assert result is not None


def test_flash_error_with_endpoint(app):
    """错误消息 - 带重定向"""
    with app.test_request_context('/test'):
        result = flash_error('操作失败', 'main.index')
        assert result is not None


def test_flash_warning_with_endpoint(app):
    """警告消息 - 带重定向"""
    with app.test_request_context('/test'):
        result = flash_warning('警告信息', 'main.index')
        assert result is not None


# ==================== localize_dt / format_local_dt 测试 ====================

def test_localize_dt_none():
    """None 值"""
    assert localize_dt(None) == ''


def test_localize_dt_naive(app):
    """无时区时间"""
    dt = datetime(2024, 1, 1, 12, 0, 0)
    result = localize_dt(dt)
    assert result is not None


def test_format_local_dt_none():
    """None 值格式化"""
    assert format_local_dt(None) == ''


def test_format_local_dt_valid(app):
    """有效时间格式化"""
    dt = datetime(2024, 1, 1, 12, 0, 0)
    result = format_local_dt(dt)
    assert '2024' in result


# ==================== get_redirect_tab 测试 ====================

def test_get_redirect_tab_draft():
    """草稿状态"""
    assert get_redirect_tab('draft') == 'draft'


def test_get_redirect_tab_confirmed():
    """已确认状态"""
    assert get_redirect_tab('confirmed') == 'confirmed'


def test_get_redirect_tab_partial():
    """部分完成状态"""
    assert get_redirect_tab('partial') == 'confirmed'


def test_get_redirect_tab_completed():
    """已完成状态"""
    assert get_redirect_tab('completed') == 'completed'


def test_get_redirect_tab_none():
    """无状态"""
    assert get_redirect_tab(None) == 'draft'


def test_get_redirect_tab_other():
    """其他状态"""
    assert get_redirect_tab('unknown') == 'draft'


# ==================== Excel 样式测试 ====================

def test_excel_constants():
    """Excel 样式常量"""
    assert EXCEL_HEADER_FONT is not None
    assert EXCEL_HEADER_FILL is not None
    assert EXCEL_THIN_BORDER is not None
    assert EXCEL_HEADER_ALIGNMENT is not None


# ==================== SystemSetting 测试 ====================

def test_system_setting_get_value_default(app, db_session):
    """获取不存在的设置 - 返回默认值"""
    result = SystemSetting.get_value('NONEXISTENT_KEY', 'default_value')
    assert result == 'default_value'


def test_system_setting_set_and_get(app, db_session):
    """设置并获取值"""
    SystemSetting.set_value('TEST_KEY', 'test_value')
    result = SystemSetting.get_value('TEST_KEY')
    assert result == 'test_value'


def test_system_setting_set_empty(app, db_session):
    """设置空值"""
    SystemSetting.set_value('EMPTY_KEY', '')
    result = SystemSetting.get_value('EMPTY_KEY')
    assert result == ''


def test_system_setting_set_none(app, db_session):
    """设置 None 值"""
    SystemSetting.set_value('NONE_KEY', None)
    result = SystemSetting.get_value('NONE_KEY')
    assert result == ''


def test_system_setting_update_existing(app, db_session):
    """更新已存在的设置"""
    SystemSetting.set_value('UPDATE_KEY', 'old_value')
    SystemSetting.set_value('UPDATE_KEY', 'new_value')
    result = SystemSetting.get_value('UPDATE_KEY')
    assert result == 'new_value'


def test_system_setting_repr(app, db_session):
    """SystemSetting __repr__"""
    setting = SystemSetting(setting_key='TEST', value='value')
    assert 'TEST' in repr(setting)


# ==================== StockLog 测试 ====================

def test_stock_log_properties(app, db_session):
    """StockLog 属性测试"""
    product = Product(code='SL_PROD', name='日志测试商品', unit='个')
    warehouse = Warehouse(code='SL_WH', name='日志测试仓库')
    db.session.add_all([product, warehouse])
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('10'),
        reference_type='stock_in',
        reference_id=1
    )
    db.session.add(log)
    db.session.commit()

    # 测试 stock_in 属性
    assert log.stock_in is None  # 没有关联的 StockIn

    # 测试 stock_out 属性
    assert log.stock_out is None  # 没有关联的 StockOut

    # 测试 reference_number 属性
    assert log.reference_number == ''  # 没有关联的单据


def test_stock_log_reference_number_purchase_order(app, db_session):
    """StockLog reference_number - 采购订单"""
    product = Product(code='SL_PROD2', name='日志测试商品2', unit='个')
    warehouse = Warehouse(code='SL_WH2', name='日志测试仓库2')
    supplier = Supplier(code='SL_SUP', name='日志测试供应商', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, supplier])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010001',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('10'),
        reference_type='purchase_order',
        reference_id=order.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'PO202401010001'


def test_stock_log_reference_number_sales_order(app, db_session):
    """StockLog reference_number - 销售订单"""
    product = Product(code='SL_PROD3', name='日志测试商品3', unit='个')
    warehouse = Warehouse(code='SL_WH3', name='日志测试仓库3')
    customer = Customer(code='SL_CUS', name='日志测试客户', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, customer])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010001',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='out',
        quantity=Decimal('5'),
        before_quantity=Decimal('10'),
        after_quantity=Decimal('5'),
        reference_type='sales_order',
        reference_id=order.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'SO202401010001'


def test_stock_log_reference_number_purchase_return(app, db_session):
    """StockLog reference_number - 采购退货"""
    product = Product(code='SL_PROD4', name='日志测试商品4', unit='个')
    warehouse = Warehouse(code='SL_WH4', name='日志测试仓库4')
    supplier = Supplier(code='SL_SUP2', name='日志测试供应商2', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, supplier])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010002',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='completed'
    )
    db.session.add(order)
    db.session.commit()

    ret = PurchaseReturn(
        return_number='PR202401010001',
        purchase_order_id=order.id,
        warehouse_id=warehouse.id,
        return_date=date.today(),
        status='completed'
    )
    db.session.add(ret)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='return_out',
        quantity=Decimal('3'),
        before_quantity=Decimal('10'),
        after_quantity=Decimal('7'),
        reference_type='purchase_return',
        reference_id=ret.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'PR202401010001'


def test_stock_log_reference_number_sales_return(app, db_session):
    """StockLog reference_number - 销售退货"""
    product = Product(code='SL_PROD5', name='日志测试商品5', unit='个')
    warehouse = Warehouse(code='SL_WH5', name='日志测试仓库5')
    customer = Customer(code='SL_CUS2', name='日志测试客户2', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, customer])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010002',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='completed'
    )
    db.session.add(order)
    db.session.commit()

    ret = SalesReturn(
        return_number='SR202401010001',
        sales_order_id=order.id,
        warehouse_id=warehouse.id,
        return_date=date.today(),
        status='completed'
    )
    db.session.add(ret)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='return_in',
        quantity=Decimal('2'),
        before_quantity=Decimal('5'),
        after_quantity=Decimal('7'),
        reference_type='sales_return',
        reference_id=ret.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'SR202401010001'


def test_stock_log_reference_number_none(app, db_session):
    """StockLog reference_number - 无引用"""
    product = Product(code='SL_PROD6', name='日志测试商品6', unit='个')
    warehouse = Warehouse(code='SL_WH6', name='日志测试仓库6')
    db.session.add_all([product, warehouse])
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='adjust_in',
        quantity=Decimal('5'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('5'),
        reference_type=None,
        reference_id=None
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == ''


def test_stock_log_reference_number_stock_in(app, db_session):
    """StockLog reference_number - 入库单"""
    product = Product(code='SL_PROD8', name='日志测试商品8', unit='个')
    warehouse = Warehouse(code='SL_WH8', name='日志测试仓库8')
    supplier = Supplier(code='SL_SUP3', name='日志测试供应商3', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, supplier])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010003',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_in = StockIn(
        receipt_number='SI202401010001',
        purchase_order_id=order.id,
        warehouse_id=warehouse.id,
        receipt_date=date.today(),
        status='completed'
    )
    db.session.add(stock_in)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('10'),
        reference_type='stock_in',
        reference_id=stock_in.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'SI202401010001'


def test_stock_log_reference_number_stock_out(app, db_session):
    """StockLog reference_number - 出库单"""
    product = Product(code='SL_PROD9', name='日志测试商品9', unit='个')
    warehouse = Warehouse(code='SL_WH9', name='日志测试仓库9')
    customer = Customer(code='SL_CUS3', name='日志测试客户3', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, customer])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010003',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_out = StockOut(
        delivery_number='OUT202401010001',
        sales_order_id=order.id,
        warehouse_id=warehouse.id,
        delivery_date=date.today(),
        status='completed'
    )
    db.session.add(stock_out)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='out',
        quantity=Decimal('5'),
        before_quantity=Decimal('10'),
        after_quantity=Decimal('5'),
        reference_type='stock_out',
        reference_id=stock_out.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.reference_number == 'OUT202401010001'


def test_stock_log_stock_in_property(app, db_session):
    """StockLog stock_in 属性"""
    product = Product(code='SL_PROD10', name='日志测试商品10', unit='个')
    warehouse = Warehouse(code='SL_WH10', name='日志测试仓库10')
    supplier = Supplier(code='SL_SUP4', name='日志测试供应商4', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, supplier])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010004',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_in = StockIn(
        receipt_number='SI202401010002',
        purchase_order_id=order.id,
        warehouse_id=warehouse.id,
        receipt_date=date.today(),
        status='completed'
    )
    db.session.add(stock_in)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('10'),
        reference_type='stock_in',
        reference_id=stock_in.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.stock_in is not None
    assert log.stock_in.id == stock_in.id


def test_stock_log_stock_out_property(app, db_session):
    """StockLog stock_out 属性"""
    product = Product(code='SL_PROD11', name='日志测试商品11', unit='个')
    warehouse = Warehouse(code='SL_WH11', name='日志测试仓库11')
    customer = Customer(code='SL_CUS4', name='日志测试客户4', contact_person='测试', phone='13800000000')
    db.session.add_all([product, warehouse, customer])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010004',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_out = StockOut(
        delivery_number='OUT202401010002',
        sales_order_id=order.id,
        warehouse_id=warehouse.id,
        delivery_date=date.today(),
        status='completed'
    )
    db.session.add(stock_out)
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='out',
        quantity=Decimal('5'),
        before_quantity=Decimal('10'),
        after_quantity=Decimal('5'),
        reference_type='stock_out',
        reference_id=stock_out.id
    )
    db.session.add(log)
    db.session.commit()

    assert log.stock_out is not None
    assert log.stock_out.id == stock_out.id


def test_stock_log_repr(app, db_session):
    """StockLog __repr__"""
    product = Product(code='SL_PROD7', name='日志测试商品7', unit='个')
    warehouse = Warehouse(code='SL_WH7', name='日志测试仓库7')
    db.session.add_all([product, warehouse])
    db.session.commit()

    log = StockLog(
        product_id=product.id,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        before_quantity=Decimal('0'),
        after_quantity=Decimal('10')
    )
    db.session.add(log)
    db.session.commit()

    assert 'in' in repr(log)


def test_receipt_repr(app, db_session):
    """Receipt __repr__"""
    customer = Customer(code='REC_CUS', name='收款测试客户', contact_person='测试', phone='13800000000')
    db.session.add(customer)
    db.session.commit()

    receipt = Receipt(
        receipt_number='REC202401010001',
        customer_id=customer.id,
        amount=Decimal('1000'),
        receipt_date=date.today(),
        payment_method='cash'
    )
    db.session.add(receipt)
    db.session.commit()

    assert 'REC202401010001' in repr(receipt)


def test_payment_repr(app, db_session):
    """Payment __repr__"""
    supplier = Supplier(code='PAY_SUP', name='付款测试供应商', contact_person='测试', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()

    payment = Payment(
        payment_number='PAY202401010001',
        supplier_id=supplier.id,
        amount=Decimal('500'),
        payment_date=date.today(),
        payment_method='bank_transfer'
    )
    db.session.add(payment)
    db.session.commit()

    assert 'PAY202401010001' in repr(payment)


def test_expense_repr(app, db_session):
    """Expense __repr__"""
    expense = Expense(
        expense_number='EXP202401010001',
        category='办公费',
        amount=Decimal('100'),
        expense_date=date.today()
    )
    db.session.add(expense)
    db.session.commit()

    assert 'EXP202401010001' in repr(expense)


def test_purchase_order_repr(app, db_session):
    """PurchaseOrder __repr__"""
    supplier = Supplier(code='PO_SUP', name='采购订单测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='PO_WH', name='采购订单测试仓库')
    db.session.add_all([supplier, warehouse])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010006',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    assert 'PO202401010006' in repr(order)


def test_purchase_order_item_repr(app, db_session):
    """PurchaseOrderItem __repr__"""
    supplier = Supplier(code='POI_SUP', name='订单明细测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='POI_WH', name='订单明细测试仓库')
    product = Product(code='POI_PROD', name='订单明细测试商品', unit='个')
    db.session.add_all([supplier, warehouse, product])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010007',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    item = PurchaseOrderItem(
        order_id=order.id,
        product_id=product.id,
        quantity=Decimal('10'),
        unit_price=Decimal('50'),
        amount=Decimal('500')
    )
    db.session.add(item)
    db.session.commit()

    assert str(item.id) in repr(item)


def test_stock_in_repr(app, db_session):
    """StockIn __repr__"""
    supplier = Supplier(code='SI_SUP', name='入库测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SI_WH', name='入库测试仓库')
    db.session.add_all([supplier, warehouse])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010008',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_in = StockIn(
        receipt_number='SI202401010003',
        purchase_order_id=order.id,
        warehouse_id=warehouse.id,
        receipt_date=date.today(),
        status='completed'
    )
    db.session.add(stock_in)
    db.session.commit()

    assert 'SI202401010003' in repr(stock_in)


def test_stock_in_item_repr(app, db_session):
    """StockInItem __repr__"""
    supplier = Supplier(code='SII_SUP', name='入库明细测试供应商', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SII_WH', name='入库明细测试仓库')
    product = Product(code='SII_PROD', name='入库明细测试商品', unit='个')
    db.session.add_all([supplier, warehouse, product])
    db.session.commit()

    order = PurchaseOrder(
        order_number='PO202401010009',
        supplier_id=supplier.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_in = StockIn(
        receipt_number='SI202401010004',
        purchase_order_id=order.id,
        warehouse_id=warehouse.id,
        receipt_date=date.today(),
        status='completed'
    )
    db.session.add(stock_in)
    db.session.commit()

    item = StockInItem(
        stock_in_id=stock_in.id,
        product_id=product.id,
        quantity=Decimal('10'),
        unit_price=Decimal('50'),
        amount=Decimal('500')
    )
    db.session.add(item)
    db.session.commit()

    assert str(item.id) in repr(item)


def test_sales_order_repr(app, db_session):
    """SalesOrder __repr__"""
    customer = Customer(code='SO_CUS', name='销售订单测试客户', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SO_WH', name='销售订单测试仓库')
    db.session.add_all([customer, warehouse])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010005',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    assert 'SO202401010005' in repr(order)


def test_sales_order_item_repr(app, db_session):
    """SalesOrderItem __repr__"""
    customer = Customer(code='SOI_CUS', name='销售明细测试客户', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SOI_WH', name='销售明细测试仓库')
    product = Product(code='SOI_PROD', name='销售明细测试商品', unit='个')
    db.session.add_all([customer, warehouse, product])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010006',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    item = SalesOrderItem(
        order_id=order.id,
        product_id=product.id,
        quantity=Decimal('10'),
        unit_price=Decimal('100'),
        amount=Decimal('1000')
    )
    db.session.add(item)
    db.session.commit()

    assert str(item.id) in repr(item)


def test_stock_out_repr(app, db_session):
    """StockOut __repr__"""
    customer = Customer(code='SO_CUS2', name='出库测试客户', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SO_WH2', name='出库测试仓库')
    db.session.add_all([customer, warehouse])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010007',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_out = StockOut(
        delivery_number='OUT202401010003',
        sales_order_id=order.id,
        warehouse_id=warehouse.id,
        delivery_date=date.today(),
        status='completed'
    )
    db.session.add(stock_out)
    db.session.commit()

    assert 'OUT202401010003' in repr(stock_out)


def test_stock_out_item_repr(app, db_session):
    """StockOutItem __repr__"""
    customer = Customer(code='SOI_CUS2', name='出库明细测试客户', contact_person='测试', phone='13800000000')
    warehouse = Warehouse(code='SOI_WH2', name='出库明细测试仓库')
    product = Product(code='SOI_PROD2', name='出库明细测试商品', unit='个')
    db.session.add_all([customer, warehouse, product])
    db.session.commit()

    order = SalesOrder(
        order_number='SO202401010008',
        customer_id=customer.id,
        warehouse_id=warehouse.id,
        order_date=date.today(),
        status='confirmed'
    )
    db.session.add(order)
    db.session.commit()

    stock_out = StockOut(
        delivery_number='OUT202401010004',
        sales_order_id=order.id,
        warehouse_id=warehouse.id,
        delivery_date=date.today(),
        status='completed'
    )
    db.session.add(stock_out)
    db.session.commit()

    item = StockOutItem(
        stock_out_id=stock_out.id,
        product_id=product.id,
        quantity=Decimal('5'),
        unit_price=Decimal('100'),
        amount=Decimal('500')
    )
    db.session.add(item)
    db.session.commit()

    assert str(item.id) in repr(item)


# ==================== create_stock_log 测试 ====================

def test_create_stock_log_in(app, db_session):
    """创建入库日志"""
    product = Product(code='CSL_PROD', name='日志测试商品', unit='个', stock_quantity=Decimal('0'))
    warehouse = Warehouse(code='CSL_WH', name='日志测试仓库')
    db.session.add_all([product, warehouse])
    db.session.commit()

    log = create_stock_log(
        product=product,
        warehouse_id=warehouse.id,
        change_type='in',
        quantity=Decimal('10'),
        reference_id=1,
        reference_type='stock_in',
        notes='测试入库',
        user_id=1
    )

    assert log is not None
    assert product.stock_quantity == Decimal('10')


def test_create_stock_log_out(app, db_session):
    """创建出库日志"""
    product = Product(code='CSL_PROD2', name='日志测试商品2', unit='个', stock_quantity=Decimal('20'))
    warehouse = Warehouse(code='CSL_WH2', name='日志测试仓库2')
    db.session.add_all([product, warehouse])
    db.session.commit()

    log = create_stock_log(
        product=product,
        warehouse_id=warehouse.id,
        change_type='out',
        quantity=Decimal('5'),
        reference_id=1,
        reference_type='stock_out',
        notes='测试出库',
        user_id=1
    )

    assert log is not None
    assert product.stock_quantity == Decimal('15')


# ==================== User 模型测试 ====================

def test_user_password_hash(app, db_session):
    """密码哈希"""
    user = User(username='hash_user', email='hash@test.com')
    user.set_password('test123')
    assert user.check_password('test123') is True
    assert user.check_password('wrong') is False


def test_user_is_authenticated(app, db_session):
    """用户认证状态"""
    user = User(username='auth_user', email='auth@test.com')
    assert user.is_authenticated is True


def test_user_default_role(app, db_session):
    """默认角色"""
    user = User(username='role_user', email='role@test.com')
    # 默认角色可能为 None 或 'user'
    assert user.role is None or user.role == 'user'


def test_user_admin_role(app, db_session):
    """管理员角色"""
    user = User(username='admin_user', email='admin@test.com', role='admin')
    assert user.role == 'admin'


def test_user_repr(app, db_session):
    """User __repr__"""
    user = User(username='repr_user', email='repr@test.com')
    assert 'repr_user' in repr(user)


def test_load_user(app, db_session):
    """加载用户"""
    from app.models.user import load_user
    user = User(username='load_user', email='load@test.com')
    user.set_password('test123')
    db.session.add(user)
    db.session.commit()

    loaded = load_user(str(user.id))
    assert loaded is not None
    assert loaded.id == user.id


def test_load_user_not_found(app, db_session):
    """加载用户 - 不存在"""
    from app.models.user import load_user
    loaded = load_user('99999')
    assert loaded is None


def test_log_repr(app, db_session):
    """Log __repr__"""
    import time
    user = User(username=f'log_user_{int(time.time())}', email='log@test.com')
    user.set_password('test123')
    db.session.add(user)
    db.session.commit()

    log = Log(
        user_id=user.id,
        action='测试操作',
        details='测试详情',
        ip_address='127.0.0.1'
    )
    db.session.add(log)
    db.session.commit()

    assert '测试操作' in repr(log)


# ==================== Product 模型测试 ====================

def test_product_creation(app, db_session):
    """商品创建"""
    product = Product(code='PROD_TEST', name='测试商品', unit='个',
                     purchase_price=50, sale_price=100)
    db.session.add(product)
    db.session.commit()
    assert product.id is not None


def test_product_stock_default(app, db_session):
    """商品默认库存"""
    product = Product(code='PROD_STOCK', name='库存测试', unit='个')
    db.session.add(product)
    db.session.commit()
    assert product.stock_quantity == 0


def test_product_repr(app, db_session):
    """Product __repr__"""
    product = Product(code='PROD_REPR', name='测试', unit='个')
    assert 'PROD_REPR' in repr(product)


# ==================== Category 模型测试 ====================

def test_category_creation(app, db_session):
    """分类创建"""
    category = Category(name='测试分类')
    db.session.add(category)
    db.session.commit()
    assert category.id is not None


def test_category_parent_child(app, db_session):
    """分类父子关系"""
    parent = Category(name='父分类')
    db.session.add(parent)
    db.session.commit()

    child = Category(name='子分类', parent_id=parent.id)
    db.session.add(child)
    db.session.commit()

    assert child.parent_id == parent.id


def test_category_repr(app, db_session):
    """Category __repr__"""
    category = Category(name='测试分类')
    assert '测试分类' in repr(category)


# ==================== Supplier 模型测试 ====================

def test_supplier_creation(app, db_session):
    """供应商创建"""
    supplier = Supplier(code='SUP_TEST', name='测试供应商', contact_person='张三', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()
    assert supplier.id is not None


def test_supplier_balance_default(app, db_session):
    """供应商默认余额"""
    supplier = Supplier(code='SUP_BAL', name='余额测试', contact_person='测试', phone='13800000000')
    db.session.add(supplier)
    db.session.commit()
    assert supplier.payable_balance == 0


def test_supplier_repr(app, db_session):
    """Supplier __repr__"""
    supplier = Supplier(code='SUP_REPR', name='测试', contact_person='测试', phone='13800000000')
    assert 'SUP_REPR' in repr(supplier)


# ==================== Customer 模型测试 ====================

def test_customer_creation(app, db_session):
    """客户创建"""
    customer = Customer(code='CUS_TEST', name='测试客户', contact_person='李四', phone='13900000000')
    db.session.add(customer)
    db.session.commit()
    assert customer.id is not None


def test_customer_balance_default(app, db_session):
    """客户默认余额"""
    customer = Customer(code='CUS_BAL', name='余额测试', contact_person='测试', phone='13900000000')
    db.session.add(customer)
    db.session.commit()
    assert customer.receivable_balance == 0


def test_customer_repr(app, db_session):
    """Customer __repr__"""
    customer = Customer(code='CUS_REPR', name='测试', contact_person='测试', phone='13900000000')
    assert 'CUS_REPR' in repr(customer)


# ==================== Warehouse 模型测试 ====================

def test_warehouse_creation(app, db_session):
    """仓库创建"""
    warehouse = Warehouse(code='WH_TEST', name='测试仓库', address='测试地址')
    db.session.add(warehouse)
    db.session.commit()
    assert warehouse.id is not None


def test_warehouse_repr(app, db_session):
    """Warehouse __repr__"""
    warehouse = Warehouse(code='WH_REPR', name='测试仓库')
    assert 'WH_REPR' in repr(warehouse)


# ==================== Constants 测试 ====================

def test_order_status_constants():
    """订单状态常量"""
    assert OrderStatus.DRAFT == 'draft'
    assert OrderStatus.CONFIRMED == 'confirmed'
    assert OrderStatus.COMPLETED == 'completed'


def test_order_status_label():
    """订单状态标签"""
    assert OrderStatus.label('draft') == '草稿'
    assert OrderStatus.label('confirmed') == '已确认'
    assert OrderStatus.label('partial') == '部分完成'
    assert OrderStatus.label('completed') == '已完成'
    assert OrderStatus.label('unknown') == 'unknown'


def test_change_type_constants():
    """变动类型常量"""
    assert ChangeType.IN == 'in'
    assert ChangeType.OUT == 'out'
    assert ChangeType.ADJUST_IN == 'adjust_in'
    assert ChangeType.ADJUST_OUT == 'adjust_out'


def test_change_type_in_out_types():
    """变动类型分类"""
    assert 'in' in ChangeType.IN_TYPES
    assert 'adjust_in' in ChangeType.IN_TYPES
    assert 'check_in' in ChangeType.IN_TYPES
    assert 'return_in' in ChangeType.IN_TYPES
    assert 'out' in ChangeType.OUT_TYPES
    assert 'adjust_out' in ChangeType.OUT_TYPES
    assert 'check_out' in ChangeType.OUT_TYPES
    assert 'return_out' in ChangeType.OUT_TYPES
    assert 'stock_transfer' in ChangeType.OUT_TYPES


def test_payment_method_constants():
    """支付方式常量"""
    assert PaymentMethod.CASH == 'cash'
    assert PaymentMethod.BANK_TRANSFER == 'bank_transfer'


def test_payment_method_label():
    """支付方式标签"""
    assert PaymentMethod.label('cash') == '现金'
    assert PaymentMethod.label('bank_transfer') == '银行转账'
    assert PaymentMethod.label('check') == '支票'
    assert PaymentMethod.label('wechat') == '微信'
    assert PaymentMethod.label('alipay') == '支付宝'
    assert PaymentMethod.label('other') == '其他'
    assert PaymentMethod.label('unknown') == 'unknown'


def test_payment_method_labels_dict():
    """支付方式标签字典"""
    assert len(PaymentMethod.LABELS) == 6
    assert 'cash' in PaymentMethod.LABELS


def test_user_role_constants():
    """用户角色常量"""
    assert UserRole.ADMIN == 'admin'
    assert UserRole.USER == 'user'


def test_valid_change_types():
    """有效的变动类型"""
    from app.constants import VALID_CHANGE_TYPES
    assert len(VALID_CHANGE_TYPES) == 9
    assert 'in' in VALID_CHANGE_TYPES
    assert 'out' in VALID_CHANGE_TYPES
    assert 'stock_transfer' in VALID_CHANGE_TYPES


# ==================== match_reference_id_filter 测试 ====================

def test_match_reference_id_filter(app, db_session):
    """匹配引用ID过滤器"""
    # 这个函数用于构建逗号分隔 reference_id 的过滤条件
    # 主要测试不抛异常
    result = match_reference_id_filter(StockLog, 1)
    assert result is not None


def test_parse_reference_ids_empty():
    """解析引用ID - 空值"""
    from app.utils_order import parse_reference_ids
    assert parse_reference_ids(None) == []
    assert parse_reference_ids('') == []


def test_parse_reference_ids_single():
    """解析引用ID - 单个ID"""
    from app.utils_order import parse_reference_ids
    assert parse_reference_ids('123') == [123]


def test_parse_reference_ids_multiple():
    """解析引用ID - 多个ID"""
    from app.utils_order import parse_reference_ids
    assert parse_reference_ids('1,2,3') == [1, 2, 3]


def test_parse_reference_ids_with_spaces():
    """解析引用ID - 带空格"""
    from app.utils_order import parse_reference_ids
    assert parse_reference_ids('1, 2, 3') == [1, 2, 3]


def test_parse_reference_ids_invalid():
    """解析引用ID - 无效值"""
    from app.utils_order import parse_reference_ids
    assert parse_reference_ids('abc') == []
    assert parse_reference_ids('1,abc,3') == []


def test_validate_order_items_empty(app, db_session):
    """验证订单明细 - 空数据"""
    from app.utils_order import validate_order_items
    assert validate_order_items([], [], []) is False


def test_validate_order_items_valid(app, db_session):
    """验证订单明细 - 有效数据"""
    from app.utils_order import validate_order_items
    assert validate_order_items(['1'], ['10'], ['50']) is True


def test_validate_order_items_invalid(app, db_session):
    """验证订单明细 - 无效数据"""
    from app.utils_order import validate_order_items
    assert validate_order_items(['abc'], ['10'], ['50']) is False
    assert validate_order_items(['0'], ['10'], ['50']) is False
    assert validate_order_items(['1'], ['0'], ['50']) is False


# ==================== rebuild_items_on_error 测试 ====================

def test_rebuild_items_on_error(app, db_session):
    """重建商品明细数据"""
    product = Product(code='RIOE_PROD', name='测试商品', unit='个',
                     purchase_price=50, sale_price=100)
    db.session.add(product)
    db.session.commit()

    product_ids = [str(product.id)]
    quantities = ['10']
    unit_prices = ['50']

    result = rebuild_items_on_error(product_ids, quantities, unit_prices)
    assert len(result) == 1
    assert result[0]['product_id'] == product.id


def test_rebuild_items_on_error_empty(app, db_session):
    """重建商品明细数据 - 空数据"""
    result = rebuild_items_on_error([], [], [])
    assert len(result) == 0


def test_rebuild_items_on_error_invalid(app, db_session):
    """重建商品明细数据 - 无效数据"""
    result = rebuild_items_on_error(['abc'], ['10'], ['50'])
    assert len(result) == 0
