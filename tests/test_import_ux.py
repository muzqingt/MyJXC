"""
导入体验测试
"""
import pytest
import io
from openpyxl import Workbook
from tests.helpers import get_page
from tests.factories import create_category


def test_import_result_page_has_statistics(app, authenticated_client, db_session):
    """导入结果页面显示统计卡片"""
    create_category(name='测试分类')

    wb = Workbook()
    ws = wb.active
    ws.append(['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存'])
    ws.append(['UX001', '测试商品', '测试分类', '个', 10.00, 20.00, 5])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/product/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    content = resp.data.decode('utf-8')
    assert 'check-circle' in content
    assert '成功' in content


def test_import_result_page_has_error_table(app, authenticated_client, db_session):
    """导入结果页面显示错误表格"""
    # 先创建一个商品，然后尝试导入重复编码
    from tests.factories import create_product
    create_product(code='DUP_UX001', name='已存在商品')

    wb = Workbook()
    ws = wb.active
    ws.append(['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存'])
    ws.append(['DUP_UX001', '重复商品', '', '个', 10.00, 20.00, 0])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/product/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    content = resp.data.decode('utf-8')
    assert '跳过' in content


def test_supplier_import_result_page(app, authenticated_client, db_session):
    """供应商导入结果页面"""
    wb = Workbook()
    ws = wb.active
    ws.append(['供应商编码', '供应商名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['UX_SUP01', '测试供应商', '张三', '13800000001', '地址', 'test@test.com'])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/suppliers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    content = resp.data.decode('utf-8')
    assert 'check-circle' in content
    assert '成功' in content


def test_customer_import_result_page(app, authenticated_client, db_session):
    """客户导入结果页面"""
    wb = Workbook()
    ws = wb.active
    ws.append(['客户编码', '客户名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['UX_CUS01', '测试客户', '李四', '13900000001', '地址', 'test@test.com'])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/customers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    content = resp.data.decode('utf-8')
    assert 'check-circle' in content
    assert '成功' in content
