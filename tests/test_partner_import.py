"""
供应商/客户导入测试
"""
import pytest
import io
from openpyxl import Workbook
from tests.helpers import get_page


# ==================== 供应商导入 ====================

def test_supplier_import_page(authenticated_client):
    """供应商导入页面可访问"""
    get_page(authenticated_client, '/partner/suppliers/import')


def test_supplier_import_template(authenticated_client):
    """供应商模板下载"""
    resp = authenticated_client.get('/partner/suppliers/import/template')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type


def test_supplier_import_success(app, authenticated_client, db_session):
    """正常导入供应商"""
    wb = Workbook()
    ws = wb.active
    ws.append(['供应商编码', '供应商名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['IMP_SUP01', '导入供应商1', '张三', '13800000001', '地址1', 'test1@test.com'])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/suppliers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '成功' in resp.data.decode('utf-8')


def test_supplier_import_duplicate(app, authenticated_client, db_session):
    """重复编码导入"""
    from tests.factories import create_supplier
    create_supplier(code='DUP_SUP01', name='已存在供应商')

    wb = Workbook()
    ws = wb.active
    ws.append(['供应商编码', '供应商名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['DUP_SUP01', '重复供应商', '', '', '', ''])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/suppliers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '跳过' in resp.data.decode('utf-8')


# ==================== 客户导入 ====================

def test_customer_import_page(authenticated_client):
    """客户导入页面可访问"""
    get_page(authenticated_client, '/partner/customers/import')


def test_customer_import_template(authenticated_client):
    """客户模板下载"""
    resp = authenticated_client.get('/partner/customers/import/template')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type


def test_customer_import_success(app, authenticated_client, db_session):
    """正常导入客户"""
    wb = Workbook()
    ws = wb.active
    ws.append(['客户编码', '客户名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['IMP_CUS01', '导入客户1', '李四', '13900000001', '地址1', 'test1@test.com'])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/customers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '成功' in resp.data.decode('utf-8')


def test_customer_import_duplicate(app, authenticated_client, db_session):
    """重复编码导入"""
    from tests.factories import create_customer
    create_customer(code='DUP_CUS01', name='已存在客户')

    wb = Workbook()
    ws = wb.active
    ws.append(['客户编码', '客户名称', '联系人', '电话', '地址', '邮箱'])
    ws.append(['DUP_CUS01', '重复客户', '', '', '', ''])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/partner/customers/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '跳过' in resp.data.decode('utf-8')
