"""
商品导入测试
"""
import pytest
import io
from openpyxl import Workbook
from tests.helpers import get_page
from tests.factories import create_category


def test_import_page(authenticated_client):
    """导入页面可访问"""
    get_page(authenticated_client, '/product/import')


def test_import_template(authenticated_client):
    """模板下载"""
    resp = authenticated_client.get('/product/import/template')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type


def test_import_success(app, authenticated_client, db_session):
    """正常导入商品"""
    create_category(name='测试分类')

    # 创建 Excel 文件
    wb = Workbook()
    ws = wb.active
    ws.append(['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存'])
    ws.append(['IMP001', '导入商品1', '测试分类', '个', 10.00, 20.00, 5])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/product/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '成功' in resp.data.decode('utf-8')


def test_import_duplicate_code(app, authenticated_client, db_session):
    """重复编码导入"""
    from tests.factories import create_product
    create_product(code='DUP001', name='已存在商品')

    # 创建包含重复编码的 Excel
    wb = Workbook()
    ws = wb.active
    ws.append(['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存'])
    ws.append(['DUP001', '重复商品', '', '个', 10.00, 20.00, 0])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/product/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '跳过' in resp.data.decode('utf-8')


def test_import_invalid_data(app, authenticated_client, db_session):
    """无效数据导入"""
    # 创建包含空编码的 Excel
    wb = Workbook()
    ws = wb.active
    ws.append(['商品编码', '商品名称', '分类名称', '单位', '采购价', '销售价', '安全库存'])
    ws.append(['', '空编码商品', '', '个', 10.00, 20.00, 0])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    resp = authenticated_client.post('/product/import', data={
        'file': (output, 'test.xlsx'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '失败' in resp.data.decode('utf-8')


def test_import_invalid_file(authenticated_client):
    """无效文件导入"""
    resp = authenticated_client.post('/product/import', data={
        'file': (io.BytesIO(b'not excel'), 'test.txt'),
    }, follow_redirects=True, content_type='multipart/form-data')

    assert resp.status_code == 200
    assert '格式' in resp.data.decode('utf-8')
